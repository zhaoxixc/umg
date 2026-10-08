import io
import time
from datetime import datetime
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Request, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy import func
from sqlalchemy.orm import Session

from openpyxl import Workbook, load_workbook

from .. import models, schemas
from ..database import get_db
from ..deps import get_current_user, record_audit, require_operator
from ..services import (add_event, build_message, get_rack_chain,
                        rack_assets_u_set, validate_position)

router = APIRouter(prefix="/api/assets", tags=["assets"])


def _asset_or_404(db: Session, asset_id: int) -> models.Asset:
    asset = db.get(models.Asset, asset_id)
    if not asset:
        raise HTTPException(404, "资产不存在")
    return asset


def _unique_code(db: Session, code: str, exclude_id: int | None = None):
    q = db.query(models.Asset).filter(models.Asset.code == code)
    if exclude_id:
        q = q.filter(models.Asset.id != exclude_id)
    return q.first()


# ---------- 查询 ----------
@router.get("", response_model=schemas.AssetPage)
def list_assets(q: str = "", status: str = "", asset_type: str = "",
                limit: int = 20, offset: int = 0,
                db: Session = Depends(get_db),
                _: models.User = Depends(get_current_user)):
    query = db.query(models.Asset)
    if status:
        query = query.filter(models.Asset.status == status)
    if asset_type:
        query = query.filter(models.Asset.asset_type == asset_type)
    if q.strip():
        kw = f"%{q.strip()}%"
        query = query.filter(
            (models.Asset.code.like(kw)) | (models.Asset.name.like(kw))
            | (models.Asset.serial.like(kw))
        )
    total = query.count()
    items = (query.order_by(models.Asset.id.desc())
                  .offset(offset).limit(min(limit, 200)).all())
    return schemas.AssetPage(total=total, items=[schemas.AssetOut.model_validate(a) for a in items])


@router.get("/overview", response_model=schemas.OverviewStats)
def overview(db: Session = Depends(get_db),
             _: models.User = Depends(get_current_user)):
    total_u = (
        db.query(func.coalesce(func.sum(models.Rack.u_count), 0)).scalar() or 0
    )
    on_rack = (
        db.query(models.Asset)
        .filter(models.Asset.status == "on_rack").all()
    )
    used_u = sum(a.u_height for a in on_rack)
    return schemas.OverviewStats(
        datacenters=db.query(models.DataCenter).count(),
        rooms=db.query(models.Room).count(),
        racks=db.query(models.Rack).count(),
        assets_on_rack=len(on_rack),
        assets_off_rack=db.query(models.Asset)
            .filter(models.Asset.status == "off_rack").count(),
        used_u=used_u,
        total_u=total_u,
    )


@router.get("/events", response_model=list[schemas.EventOut])
def all_events(scope: str = "realtime", limit: int = 50, offset: int = 0,
               db: Session = Depends(get_db),
               _: models.User = Depends(get_current_user)):
    q = db.query(models.AssetEvent).order_by(
        models.AssetEvent.created_at.desc(), models.AssetEvent.id.desc())
    if scope == "realtime":
        return q.limit(20).all()
    return q.offset(offset).limit(min(limit, 200)).all()


# ---------- 新增（可同时上架） ----------
@router.post("", response_model=schemas.AssetOut)
def create_asset(body: schemas.AssetIn, request: Request,
                 db: Session = Depends(get_db),
                 actor: models.User = Depends(require_operator)):
    if _unique_code(db, body.code):
        raise HTTPException(400, f"资产编号 {body.code} 已存在")
    if body.u_height < 1:
        raise HTTPException(400, "资产高度至少 1U")
    data = body.model_dump()
    rack = None
    if data["rack_id"] is not None:
        rack = db.get(models.Rack, data["rack_id"])
        if not rack:
            raise HTTPException(404, "目标机柜不存在")
        if not data["u_start"]:
            raise HTTPException(400, "上架必须指定起始 U 位")
        try:
            validate_position(db, rack, data["u_start"], data["u_height"])
        except ValueError as e:
            raise HTTPException(400, str(e))
        data["status"] = "on_rack"
    else:
        data["status"] = "off_rack"
        data["u_start"] = None
    asset = models.Asset(**data)
    db.add(asset)
    db.flush()
    record_audit(db, request, actor, "asset_create",
                 f"新增资产:{asset.code}({asset.name})")
    if rack:
        add_event(db, asset, "rack_on", build_message(db, rack, asset, "上架"))
    db.commit()
    db.refresh(asset)
    return asset


# ---------- 编辑基本信息 ----------
@router.put("/{asset_id}", response_model=schemas.AssetOut)
def update_asset(asset_id: int, body: schemas.AssetUpdate, request: Request,
                 db: Session = Depends(get_db),
                 actor: models.User = Depends(require_operator)):
    asset = _asset_or_404(db, asset_id)
    for field, value in body.model_dump(exclude_unset=True).items():
        setattr(asset, field, value)
    if asset.status == "on_rack" and asset.u_height > 1:
        occupied = rack_assets_u_set(db, asset.rack_id, exclude_asset_id=asset.id)
        for u in range(asset.u_start + 1, asset.u_start + asset.u_height):
            if u in occupied:
                raise HTTPException(400, f"调整高度失败：U 位 {u}U 已被占用")
    record_audit(db, request, actor, "asset_update", f"编辑资产:{asset.code}")
    db.commit()
    db.refresh(asset)
    return asset


# ---------- 上架 / 移位 ----------
@router.post("/{asset_id}/rack-on", response_model=schemas.AssetOut)
def rack_on(asset_id: int, body: schemas.RackOnIn, request: Request,
            db: Session = Depends(get_db),
            actor: models.User = Depends(require_operator)):
    asset = _asset_or_404(db, asset_id)
    rack = db.get(models.Rack, body.rack_id)
    if not rack:
        raise HTTPException(404, "目标机柜不存在")
    try:
        validate_position(db, rack, body.u_start, asset.u_height,
                          exclude_asset_id=asset.id)
    except ValueError as e:
        raise HTTPException(400, str(e))
    moving = asset.status == "on_rack"
    old_rack_id = asset.rack_id
    asset.rack_id = rack.id
    asset.u_start = body.u_start
    asset.status = "on_rack"
    action = "移位" if moving else "上架"
    if moving and old_rack_id != rack.id:
        action = "移位"
    add_event(db, asset, "move" if action == "移位" else "rack_on",
              build_message(db, rack, asset, action))
    record_audit(db, request, actor, "asset_" + ("move" if action == "移位" else "rack_on"),
                 f"{action}:资产:{asset.code} -> 机柜:{rack.name} {body.u_start}U")
    db.commit()
    db.refresh(asset)
    return asset


# ---------- 下架 ----------
@router.post("/{asset_id}/rack-off", response_model=schemas.AssetOut)
def rack_off(asset_id: int, request: Request, db: Session = Depends(get_db),
             actor: models.User = Depends(require_operator)):
    asset = _asset_or_404(db, asset_id)
    if asset.status != "on_rack":
        raise HTTPException(400, "资产已在架下状态")
    rack = db.get(models.Rack, asset.rack_id)
    dc, room, _ = get_rack_chain(db, rack)
    message = (
        f"数据中心:{dc}|机房:{room}|机柜:{rack.name}|"
        f"资产编号:{asset.code},资产序列号:{asset.serial}|U位:{asset.u_start}U|下架"
    )
    asset.status = "off_rack"
    asset.rack_id = None
    asset.u_start = None
    add_event(db, asset, "rack_off", message)
    record_audit(db, request, actor, "asset_rack_off", f"下架资产:{asset.code}")
    db.commit()
    db.refresh(asset)
    return asset


# ---------- 删除 ----------
@router.delete("/{asset_id}", response_model=schemas.MessageOut)
def delete_asset(asset_id: int, request: Request, db: Session = Depends(get_db),
                 actor: models.User = Depends(require_operator)):
    asset = _asset_or_404(db, asset_id)
    if asset.status == "on_rack":
        raise HTTPException(400, "在架资产不能删除，请先下架")
    code = asset.code
    db.delete(asset)
    record_audit(db, request, actor, "asset_delete", f"删除资产:{code}")
    db.commit()
    return {"message": "已删除"}


# ---------- 单资产轨迹 ----------
@router.get("/{asset_id}/history", response_model=list[schemas.AssetHistoryOut])
def asset_history(asset_id: int, limit: int = 100, db: Session = Depends(get_db),
                  _: models.User = Depends(get_current_user)):
    asset = _asset_or_404(db, asset_id)
    return (
        db.query(models.AssetEvent)
        .filter(models.AssetEvent.asset_id == asset.id)
        .order_by(models.AssetEvent.created_at.desc(), models.AssetEvent.id.desc())
        .limit(min(limit, 200))
        .all()
    )


# ---------- Excel 导出 / 导入 ----------
# 与前端「设备信息」面板字段保持同步
EXPORT_HEADERS = ["资产编号", "资产名称", "资产序列号", "资产类型",
                  "资产品牌", "资产型号", "资产高度(U)", "所在机柜", "起始U位",
                  "业务IP地址", "BMC管理地址", "备注"]

ASSET_TYPE_OPTIONS = ["服务器", "交换机", "路由器", "存储", "防火墙", "其他"]


def _export_filename() -> str:
    return f"assets_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"


@router.get("/export")
def export_assets(q: str = "", status: str = "",
                  db: Session = Depends(get_db),
                  _: models.User = Depends(get_current_user)):
    query = db.query(models.Asset)
    if status:
        query = query.filter(models.Asset.status == status)
    if q.strip():
        kw = f"%{q.strip()}%"
        query = query.filter(
            (models.Asset.code.like(kw)) | (models.Asset.name.like(kw))
            | (models.Asset.serial.like(kw))
        )
    assets = query.order_by(models.Asset.id).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "资产清单"
    ws.append(EXPORT_HEADERS)
    for a in assets:
        rack = db.get(models.Rack, a.rack_id) if a.rack_id else None
        ws.append([
            a.code, a.name, a.serial, a.asset_type,
            a.brand, a.model, a.u_height,
            rack.name if rack else "",
            a.u_start or "",
            a.ip_address, a.bmc_address, a.remark,
        ])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = _export_filename()
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


@router.get("/import-template")
def import_template(_: models.User = Depends(get_current_user)):
    wb = Workbook()
    ws = wb.active
    ws.title = "资产导入"
    ws.append(EXPORT_HEADERS)
    ws.append(["IMP0001", "示例服务器", "SN00001", "服务器",
               "华为", "2288H V5", 2, "", "",
               "192.168.1.101", "192.168.100.101", ""])
    ws.append(["IMP0002", "示例交换机", "SN00002", "交换机",
               "华为", "S5720-28P-LI-AC", 1, "CPJF-S-24P", 36,
               "192.168.1.102", "", ""])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote('asset_import_template.xlsx')}"},
    )


def _cell_str(v) -> str:
    return str(v).strip() if v is not None else ""


@router.post("/import", response_model=schemas.ImportResult)
async def import_assets(file: UploadFile, request: Request,
                        db: Session = Depends(get_db),
                        actor: models.User = Depends(require_operator)):
    try:
        wb = load_workbook(io.BytesIO(await file.read()), read_only=True)
    except Exception:
        raise HTTPException(400, "无法解析文件，请使用系统提供的导入模板")
    ws = wb.active

    racks = {r.name: r for r in db.query(models.Rack).all()}
    exist_codes = {c for (c,) in db.query(models.Asset.code).all()}
    all_types = set(ASSET_TYPE_OPTIONS) | {
        t for (t,) in db.query(models.Asset.asset_type).distinct()}

    success, failed = 0, []
    for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(v is not None and _cell_str(v) for v in row):
            continue
        vals = [_cell_str(v) if i != 6 else v for i, v in enumerate(row)]
        (code, name, serial, asset_type,
         brand, model, u_height, rack_name, u_start,
         ip_address, bmc_address, remark) = (vals + [""] * 12)[:12]
        try:
            # 每行使用独立 SAVEPOINT，单行失败不影响其它行
            with db.begin_nested():
                if not code:
                    raise ValueError("资产编号不能为空")
                if not name:
                    raise ValueError("资产名称不能为空")
                if code in exist_codes:
                    raise ValueError(f"资产编号 {code} 已存在")
                try:
                    u_height = int(u_height) if u_height not in (None, "") else 1
                except (TypeError, ValueError):
                    raise ValueError(f"资产高度「{u_height}」无效")
                if not 1 <= u_height <= 10:
                    raise ValueError("资产高度需在 1-10U 之间")
                asset_type = asset_type or "服务器"
                if asset_type not in all_types:
                    asset_type = "其他"
                rack = None
                u_start_i = None
                if rack_name:
                    rack = racks.get(rack_name)
                    if not rack:
                        raise ValueError(f"机柜「{rack_name}」不存在")
                    try:
                        u_start_i = int(u_start)
                    except (TypeError, ValueError):
                        raise ValueError("填写了机柜时必须填写有效的起始U位")
                    validate_position(db, rack, u_start_i, u_height)
                asset = models.Asset(
                    code=code, name=name, serial=serial,
                    asset_type=asset_type, brand=brand, model=model,
                    ip_address=ip_address, bmc_address=bmc_address,
                    remark=remark[:255], u_height=u_height,
                    rack_id=rack.id if rack and u_start_i else None,
                    u_start=u_start_i if rack and u_start_i else None,
                    status="on_rack" if rack and u_start_i else "off_rack",
                )
                db.add(asset)
                db.flush()
            exist_codes.add(code)
            if rack and u_start_i:
                add_event(db, asset, "rack_on",
                          build_message(db, rack, asset, "上架"))
            success += 1
        except ValueError as e:
            failed.append({"row": row_no, "error": str(e)})
        except Exception as e:  # noqa: BLE001
            failed.append({"row": row_no, "error": f"数据异常：{e}"})
    record_audit(db, request, actor, "asset_import",
                 f"导入资产:成功{success}条,失败{len(failed)}条")
    db.commit()
    return schemas.ImportResult(success=success, failed=failed)
